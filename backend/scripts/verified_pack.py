"""Đọc EVFTA_20_ma_da_xac_minh.xlsx thành dòng dữ liệu. Hàm thuần: không đụng DB.

Quy đổi RoO (đã chốt với người dùng): dòng 'All fish and crustaceans… wholly obtained' không kèm
điều kiện tàu → WO, không cần chuyên gia. Mọi dòng khác (nguyên liệu Chương X thuần tuý, dung sai,
ngưỡng đường, điều kiện tàu) chưa có loại quy tắc tương ứng → WO + requires_expert=true (máy tính
luôn 'chưa kết luận'), nguyên văn Annex II nằm ở rule_text. Không đoán, không tự duyệt.
"""

import datetime as dt
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl

from app.modules.catalog.schemas import HsCodeIn
from app.modules.compliance.admin_schemas import CountryTermIn, RooRuleIn, TariffLineIn
from app.modules.compliance.models import DutyType, RuleType

DEFAULT_XLSX = (
    Path(__file__).resolve().parents[1] / "data" / "compliance" / "EVFTA_20_ma_da_xac_minh.xlsx"
)
SHEET_CODES, SHEET_VAT, SHEET_PARAMS = "20 ma da xac minh", "3 nuoc DE-FR-NL", "Tham so"
# Ngày thuế EVFTA về 0% theo bảng lộ trình ở sheet 'Tham so' (A: 1/8/2020, B3: 1/1/2023 …).
ZERO_FROM = {
    "A": dt.date(2020, 8, 1),
    "B3": dt.date(2023, 1, 1),
    "B5": dt.date(2025, 1, 1),
    "B7": dt.date(2027, 1, 1),
}
RULE_SOURCE = "EVFTA Nghị định thư 1, Annex II (file EVFTA_20_ma_da_xac_minh.xlsx)"
VAT_SOURCE = (
    "File EVFTA_20_ma_da_xac_minh.xlsx, sheet 'Tham so' (VAT thực phẩm); nguồn thứ cấp, "
    "cần đối chiếu TEDB của Uỷ ban châu Âu trước go-live"
)
VAT_COLUMNS = {"DE": ("VAT DE", "de"), "FR": ("VAT FR", "fr"), "NL": ("VAT NL", "nl")}
CATEGORY = {"Thuỷ sản": "seafood", "Rau, nấm": "agriculture", "Trái cây": "agriculture",
            "Tham chiếu 0%": "agriculture"}  # fmt: skip
# Tên tiếng Anh chỉ để tìm kiếm trong danh mục, không phải câu chữ pháp lý.
EN_NAMES = {
    "03061792": "Frozen shrimps and prawns, genus Penaeus",
    "03061799": "Frozen shrimps and prawns, other",
    "03046200": "Frozen fillets of catfish (Pangasius spp.)",
    "03043200": "Fresh or chilled fillets of catfish (Pangasius spp.)",
    "03032400": "Frozen catfish (Pangasius spp.), whole",
    "03048700": "Frozen fillets of tunas",
    "03034290": "Frozen yellowfin tunas, whole",
    "03077100": "Live, fresh or chilled clams, cockles and ark shells",
    "07123200": "Dried wood ears (Auricularia spp.)",
    "07123900": "Other dried mushrooms and truffles",
    "07108069": "Frozen mushrooms, other",
    "07108095": "Frozen vegetables, other",
    "07096099": "Fresh peppers of the genus Capsicum or Pimenta, other",
    "08106000": "Durians, fresh",
    "08109075": "Other fresh fruit (longan, rambutan and similar)",
    "08119085": "Frozen tropical fruit, without added sugar",
    "08043000": "Pineapples, fresh or dried",
    "08055090": "Limes, fresh or dried",
    "08013200": "Cashew nuts, shelled",
    "08109020": "Fresh pitahaya, lychees, passion fruit, jackfruit, carambola",
}  # fmt: skip
WHOLLY_OBTAINED = "All fish and crustaceans"
VESSEL_MARKER = "điều kiện tàu"


@dataclass(frozen=True)
class PackRow:
    cn_code: str
    group: str
    name_vi: str
    name_en: str
    category: str
    mfn_rate: str  # %
    evfta_rate: str  # %
    staging: str
    zero_from: dt.date
    rule_text: str
    requires_expert: bool
    note: str | None


@dataclass(frozen=True)
class TermRow:
    cn_code: str
    country: str
    vat_rate: str  # %
    label_languages: str
    note: str | None


@dataclass(frozen=True)
class Pack:
    year: int
    rows: list[PackRow]
    terms: list[TermRow]


def _pct(value: Any) -> str:
    """0.055 → '5.5', 0 → '0' (Decimal từ chuỗi, không qua phép tính float)."""
    return format((Decimal(str(value)) * 100).normalize(), "f")


def _text(value: Any) -> str | None:
    return str(value).strip() or None if value is not None else None


def _table(ws: Any) -> list[dict[str, Any]]:
    header = [
        str(c).strip() if c is not None else ""
        for c in next(ws.iter_rows(max_row=1, values_only=True))
    ]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values, strict=False))
        if row.get(header[1]) is not None and str(values[0] or "").strip().isdigit():
            rows.append(row)
    return rows


def _year(ws: Any) -> int:
    for label, value, *_ in ws.iter_rows(values_only=True):
        if label == "Năm tính thuế EVFTA":
            return int(value)
    raise ValueError("thiếu 'Năm tính thuế EVFTA' ở sheet Tham so")


def load_pack(path: Path = DEFAULT_XLSX) -> Pack:
    wb = openpyxl.load_workbook(path, data_only=True)
    year = _year(wb[SHEET_PARAMS])
    vat_by_code = {str(r["Mã CN"]): r for r in _table(wb[SHEET_VAT])}
    rows: list[PackRow] = []
    terms: list[TermRow] = []
    for r in _table(wb[SHEET_CODES]):
        code = str(r["Mã CN (biểu EVFTA)"])
        staging = str(r["Nhóm lộ trình EVFTA"]).strip()
        if staging not in ZERO_FROM:
            raise ValueError(f"{code}: nhóm lộ trình lạ '{staging}'")
        rule_text = str(r["Quy tắc xuất xứ – nguyên văn Annex II"]).strip()
        logic = str(r["Loại logic cho máy tính RoO"])
        wholly = rule_text.startswith(WHOLLY_OBTAINED) and VESSEL_MARKER not in logic
        rows.append(
            PackRow(
                cn_code=code,
                group=str(r["Nhóm"]).strip(),
                name_vi=str(r["Mô tả"]).strip(),
                name_en=EN_NAMES[code],
                category=CATEGORY[str(r["Nhóm"]).strip()],
                mfn_rate=_pct(r["Thuế cơ sở / MFN tham chiếu"]),
                evfta_rate=_pct(r["Thuế EVFTA năm tính"]),
                staging=staging,
                zero_from=ZERO_FROM[staging],
                rule_text=rule_text,
                requires_expert=not wholly,
                note=_text(r.get("Ghi chú")),
            )
        )
        vat = vat_by_code[code]
        for country, (column, language) in VAT_COLUMNS.items():
            terms.append(
                TermRow(
                    code,
                    country,
                    _pct(vat[column]),
                    language,
                    _text(vat.get("Khác biệt quốc gia cần lưu ý")),
                )
            )
    return Pack(year, rows, terms)


def hs_rows(pack: Pack) -> list[HsCodeIn]:
    return [
        HsCodeIn(
            code=r.cn_code,
            name_vi=r.name_vi[:255],
            name_en=r.name_en,
            category=r.category,
            is_calculator_supported=True,
        )
        for r in pack.rows
    ]


def tariff_rows(pack: Pack) -> list[tuple[int, TariffLineIn]]:
    return [
        (
            n,
            TariffLineIn(
                hs_code=r.cn_code,
                destination="EU",
                duty_type=DutyType.ad_valorem,
                mfn_rate=r.mfn_rate,
                evfta_rate_current=r.evfta_rate,
                staging_category=r.staging,
                zero_from=r.zero_from,
                quota_required=False,
                condition_note=r.note,
                valid_from=r.zero_from,
            ),
        )
        for n, r in enumerate(pack.rows, start=2)
    ]


def psr_rows(pack: Pack) -> list[tuple[int, RooRuleIn]]:
    return [
        (
            n,
            RooRuleIn(
                hs_code=r.cn_code,
                rule_type=RuleType.WO,
                rule_text=r.rule_text,
                requires_expert=r.requires_expert,
                source=RULE_SOURCE,
                valid_from=dt.date(2020, 8, 1),
            ),
        )
        for n, r in enumerate(pack.rows, start=2)
    ]


def term_rows(pack: Pack) -> list[tuple[int, CountryTermIn]]:
    return [
        (
            n,
            CountryTermIn(
                hs_code=t.cn_code,
                country=t.country,
                vat_rate=t.vat_rate,
                label_languages=t.label_languages,
                note=t.note,
                source=VAT_SOURCE,
                valid_from=dt.date(pack.year, 1, 1),
            ),
        )
        for n, t in enumerate(pack.terms, start=2)
    ]
