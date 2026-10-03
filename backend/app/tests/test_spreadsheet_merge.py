"""Gộp ô theo nhóm khi xuất Excel (merge center)."""

import io

from openpyxl import load_workbook

from app.core.spreadsheet import Column, build_workbook

COLS = (Column("loai"), Column("gia_tri"), Column("ten"))


def merged(rows: list[dict[str, str]]) -> set[str]:
    ws = load_workbook(io.BytesIO(build_workbook(COLS, rows, merge=("loai", "gia_tri"))))["data"]
    return {str(r) for r in ws.merged_cells.ranges}


def test_merges_each_group_vertically_and_centers() -> None:
    rows = [
        {"loai": "Tên miền", "gia_tri": "a.vn", "ten": "A"},
        {"loai": "Tên miền", "gia_tri": "a.vn", "ten": "B"},
        {"loai": "Tên miền", "gia_tri": "b.vn", "ten": "C"},  # cùng loại, khác giá trị → nhóm mới
        {"loai": "Tên miền", "gia_tri": "b.vn", "ten": "D"},
        {"loai": "Mã số thuế", "gia_tri": "1", "ten": "E"},  # nhóm một dòng → không gộp
    ]
    assert merged(rows) == {"A2:A3", "B2:B3", "A4:A5", "B4:B5"}
    ws = load_workbook(io.BytesIO(build_workbook(COLS, rows, merge=("loai", "gia_tri"))))["data"]
    assert ws["A2"].alignment.horizontal == "center"
    assert ws["A2"].alignment.vertical == "center"
    assert [ws.cell(row=r, column=3).value for r in range(2, 7)] == ["A", "B", "C", "D", "E"]


def test_no_merge_by_default() -> None:
    rows = [{"loai": "x", "gia_tri": "1", "ten": "A"}, {"loai": "x", "gia_tri": "1", "ten": "B"}]
    ws = load_workbook(io.BytesIO(build_workbook(COLS, rows)))["data"]
    assert not ws.merged_cells.ranges
