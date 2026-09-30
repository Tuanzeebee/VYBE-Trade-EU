"""Chống mạo danh (I11): danh sách chặn, kiểm danh tính, cụm tài khoản, hàng đợi. Dữ liệu SYNTHETIC.

"Xong khi": định danh bị chặn không đăng ký lại được; chưa chứng minh quyền sở hữu thì không lên
evfta_verified; tín hiệu chỉ đẩy lên hàng đợi, không đổi trạng thái xác minh.
"""

import datetime as dt
import hashlib
import io
import uuid

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.audit import AuditLog
from app.core.errors import AppError
from app.core.spreadsheet import XLSX
from app.modules.companies import service as companies
from app.modules.companies.tests.helpers import PASSWORD, company_body, login_as, product_body
from app.modules.verification import evidence_service
from app.modules.verification.identity import hash_name
from app.modules.verification.models import (
    ApprovalStatus,
    BlocklistIdentifier,
    Evidence,
    EvidenceCheck,
    IdentifierType,
)
from app.modules.verification.tests.helpers import TODAY, add_rule, add_type, body
from app.modules.verification.tests.test_verification_requests import (
    QUEUE,
    SUBMIT,
    login_admin,
    new_exporter,
)

BLOCK = "/api/admin/blocklist"


def identity_url(company_id: str | uuid.UUID) -> str:
    return f"/api/admin/companies/{company_id}/identity"


def checks_url(company_id: str | uuid.UUID) -> str:
    return f"/api/admin/companies/{company_id}/identity-checks"


async def block(session: AsyncSession, admin_id: uuid.UUID, kind: str, value: str) -> None:
    session.add(
        BlocklistIdentifier(
            identifier_type=IdentifierType(kind),
            value=value,
            reason="đã xác nhận",
            added_by=admin_id,
        )
    )
    await session.flush()


def error_code(r: object) -> str:
    return str(r.json()["error"]["code"])  # type: ignore[attr-defined]


# ── Danh sách chặn: không đăng ký / tạo hồ sơ / nộp bằng chứng lại được ─────────────────────
async def test_blocked_email_domain_cannot_register(
    api_client: AsyncClient, db_session: AsyncSession, reviewer_id: uuid.UUID
) -> None:
    await block(db_session, reviewer_id, "domain", "fraud-co.vn")
    r = await api_client.post(
        "/api/auth/register",
        json={
            "email": "sales@Fraud-Co.vn",
            "password": PASSWORD,
            "role": "exporter",
            "accept_terms": True,
        },
    )
    assert r.status_code == 403 and error_code(r) == "identifier_blocked"


async def test_blocked_phone_cannot_register_in_any_format(
    api_client: AsyncClient, db_session: AsyncSession, reviewer_id: uuid.UUID
) -> None:
    await block(db_session, reviewer_id, "phone", "84912345678")
    r = await api_client.post(
        "/api/auth/register",
        json={
            "email": "new@ok.vn",
            "password": PASSWORD,
            "role": "exporter",
            "phone": "0912 345 678",
            "accept_terms": True,
        },
    )
    assert r.status_code == 403 and error_code(r) == "identifier_blocked"


async def test_blocked_tax_id_cannot_create_company(
    api_client: AsyncClient, db_session: AsyncSession, reviewer_id: uuid.UUID
) -> None:
    await block(db_session, reviewer_id, "tax_id", "0314892345")
    await login_as(api_client, "exporter", "exp@x.vn")
    r = await api_client.post("/api/me/company", json=company_body(tax_id="0314-892-345"))
    assert r.status_code == 403 and error_code(r) == "identifier_blocked"


async def test_blocked_website_cannot_be_set_on_update(
    api_client: AsyncClient, db_session: AsyncSession, company_id: uuid.UUID, reviewer_id: uuid.UUID
) -> None:
    await block(db_session, reviewer_id, "domain", "shell-co.vn")
    r = await api_client.patch("/api/me/company", json={"website": "https://www.shell-co.vn"})
    assert r.status_code == 403 and error_code(r) == "identifier_blocked"


async def test_evidence_stores_file_hash_and_blocked_hash_is_refused(
    api_client: AsyncClient, db_session: AsyncSession, company_id: uuid.UUID, reviewer_id: uuid.UUID
) -> None:
    await add_type(db_session, reviewer_id)
    ok = body(str(company_id))
    r = await api_client.post("/api/exporter/evidences", json=ok)
    assert r.status_code == 201, r.text
    row = await db_session.get(Evidence, uuid.UUID(r.json()["id"]))
    assert row is not None
    # FakeStorage trả nội dung "fake:<key>" cho file chưa ghi thật.
    assert row.file_sha256 == hashlib.sha256(b"fake:" + ok["file_key"].encode()).hexdigest()

    bad = body(str(company_id))
    digest = hashlib.sha256(b"fake:" + bad["file_key"].encode()).hexdigest()
    await block(db_session, reviewer_id, "file_sha256", digest)
    r = await api_client.post("/api/exporter/evidences", json=bad)
    assert r.status_code == 403 and error_code(r) == "identifier_blocked"


class _EmptyStorage:
    async def get(self, key: str) -> bytes | None:
        return None


async def test_evidence_file_must_exist(db_session: AsyncSession) -> None:
    with pytest.raises(AppError) as exc:
        await evidence_service._file_hash(db_session, _EmptyStorage(), "evidence/x/a.pdf")  # type: ignore[arg-type]
    assert (exc.value.status_code, exc.value.code) == (422, "file_not_uploaded")


# ── Quản lý danh sách chặn (admin, có audit) ───────────────────────────────────────────────
async def test_admin_blocklist_normalizes_rejects_duplicates_and_audits(
    api_client: AsyncClient, db_session: AsyncSession
) -> None:
    await login_admin(api_client, db_session)
    r = await api_client.post(
        BLOCK, json={"identifier_type": "phone", "value": "0912 345 678", "reason": "mạo danh"}
    )
    assert r.status_code == 201, r.text
    assert r.json()["value"] == "84912345678"
    dup = await api_client.post(
        BLOCK, json={"identifier_type": "phone", "value": "+84912345678", "reason": "x"}
    )
    assert dup.status_code == 409
    free = await api_client.post(
        BLOCK, json={"identifier_type": "domain", "value": "gmail.com", "reason": "x"}
    )
    assert free.status_code == 422
    assert [e["value"] for e in (await api_client.get(BLOCK)).json()] == ["84912345678"]
    assert (await api_client.delete(f"{BLOCK}/{r.json()['id']}")).status_code == 204
    actions = list(await db_session.scalars(select(AuditLog.action_type)))
    assert "blocklist.add" in actions and "blocklist.remove" in actions


@pytest.mark.parametrize(
    ("method", "path"),
    [
        ("GET", "/api/admin/companies/{id}/identity"),
        ("POST", "/api/admin/companies/{id}/identity-checks"),
        ("GET", "/api/admin/identity-clusters"),
        ("GET", "/api/admin/identity-clusters/export.xlsx"),
        ("GET", "/api/admin/blocklist"),
        ("POST", "/api/admin/blocklist"),
        ("DELETE", "/api/admin/blocklist/{id}"),
    ],
)
async def test_identity_routes_need_admin(api_client: AsyncClient, method: str, path: str) -> None:
    url = path.format(id=uuid.uuid4())
    payload = {"check_type": "phone_callback", "result": "match"} if method == "POST" else None
    assert (await api_client.request(method, url, json=payload)).status_code == 401
    await login_as(api_client, "exporter", "exp@x.vn")
    assert (await api_client.request(method, url, json=payload)).status_code == 403


# ── Quyền sở hữu là điều kiện của evfta_verified; ghi kiểm không đổi trạng thái ───────────────
@pytest.mark.usefixtures("hs_seeded")
async def test_ownership_check_gates_evfta_verified(
    api_client: AsyncClient, db_session: AsyncSession, company_id: uuid.UUID, reviewer_id: uuid.UUID
) -> None:
    await add_type(db_session, reviewer_id)
    await add_rule(db_session, reviewer_id, "iso_9001")
    await api_client.post("/api/exporter/products", json=product_body(hs_code="090121"))
    db_session.add(
        Evidence(
            company_id=company_id,
            type_code="iso_9001",
            file_key=f"evidence/{company_id}/a.pdf",
            issued_at=TODAY - dt.timedelta(days=3),
            approval_status=ApprovalStatus.approved,
        )
    )
    now = dt.datetime.now(dt.UTC)
    await companies.set_verification_state(
        db_session, company_id, status="verified", level="basic", verified_at=now, expires_at=None
    )
    await evidence_service.sync_level(db_session, company_id)
    state = await companies.get_verification_state(db_session, company_id)
    assert state.level == "basic"  # đủ bằng chứng nhưng chưa chứng minh quyền sở hữu

    await login_admin(api_client, db_session)
    r = await api_client.post(
        checks_url(company_id), json={"check_type": "phone_callback", "result": "match"}
    )
    assert r.status_code == 201, r.text
    state = await companies.get_verification_state(db_session, company_id)
    assert (state.status, state.level) == ("verified", "evfta_verified")
    assert (await api_client.get(identity_url(company_id))).json()["ownership_proven"] is True

    await api_client.post(
        checks_url(company_id),
        json={"check_type": "phone_callback", "result": "mismatch", "note": "số không nghe"},
    )
    state = await companies.get_verification_state(db_session, company_id)
    assert (state.status, state.level) == ("verified", "basic")
    actions = list(await db_session.scalars(select(AuditLog.action_type)))
    assert actions.count("identity_check.create") == 2


async def test_registry_facts_are_hashed_and_only_flag(
    api_client: AsyncClient, db_session: AsyncSession, company_id: uuid.UUID
) -> None:
    await login_admin(api_client, db_session)
    r = await api_client.post(
        checks_url(company_id),
        json={
            "check_type": "registry_lookup",
            "result": "mismatch",
            "registry": {
                "legal_representative": "Nguyễn Văn An",
                "founded_year": 2025,
                "tax_status": "inactive",
            },
        },
    )
    assert r.status_code == 201, r.text
    facts = r.json()["facts"]
    assert facts["legal_representative_hash"] == hash_name("nguyen van an")
    assert "Nguyễn" not in str(facts)  # không lưu tên người
    identity = (await api_client.get(identity_url(company_id))).json()
    codes = {s["code"] for s in identity["signals"]}
    assert {"tax_inactive", "founded_mismatch"} <= codes  # company_body khai founded_year 2018
    state = await companies.get_verification_state(db_session, company_id)
    assert state.status == "unverified"  # tín hiệu không bao giờ đổi trạng thái


async def test_registry_facts_only_for_registry_lookup(
    api_client: AsyncClient, db_session: AsyncSession, company_id: uuid.UUID
) -> None:
    await login_admin(api_client, db_session)
    r = await api_client.post(
        checks_url(company_id),
        json={
            "check_type": "phone_callback",
            "result": "match",
            "registry": {"founded_year": 2020},
        },
    )
    assert r.status_code == 422


async def test_unknown_company_is_404(api_client: AsyncClient, db_session: AsyncSession) -> None:
    await login_admin(api_client, db_session)
    assert (await api_client.get(identity_url(uuid.uuid4()))).status_code == 404


# ── Gom cụm và hàng đợi: cờ chỉ đẩy lên đầu ───────────────────────────────────────────────
async def test_clusters_and_flagged_requests_go_first(
    api_client: AsyncClient, db_session: AsyncSession
) -> None:
    clean = company_body(
        legal_name="Công ty C",
        tax_id="0100000001",
        website="https://c-corp.vn",
        contact_email="info@c-corp.vn",
    )
    await api_client.post("/api/auth/logout")
    await login_as(api_client, "exporter", "owner@c-corp.vn")
    c_id = (await api_client.post("/api/me/company", json=clean)).json()["id"]
    await api_client.post(SUBMIT)  # C nộp trước
    a_id = await new_exporter(api_client, "a@x-one.vn", "Công ty A")  # cùng MST mặc định với B
    await api_client.post(SUBMIT)
    b_id = await new_exporter(api_client, "b@x-two.vn", "Công ty B")
    await api_client.post(SUBMIT)

    await login_admin(api_client, db_session)
    queue = (await api_client.get(QUEUE)).json()
    assert [q["company_id"] for q in queue] == [a_id, b_id, c_id]
    assert {"code": "shared_tax_id", "severity": "high"} in queue[0]["signals"]
    assert queue[2]["signals"] == []
    for cid in (a_id, b_id, c_id):  # cờ không đổi trạng thái
        state = await companies.get_verification_state(db_session, uuid.UUID(cid))
        assert state.status == "pending"

    clusters = (await api_client.get("/api/admin/identity-clusters")).json()
    tax = [c for c in clusters if c["identifier_type"] == "tax_id"]
    assert len(tax) == 1
    assert {x["id"] for x in tax[0]["companies"]} == {a_id, b_id}

    # Xuất Excel: mỗi doanh nghiệp một dòng, định danh + giá trị gộp ô theo cụm.
    r = await api_client.get("/api/admin/identity-clusters/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(XLSX)
    ws = load_workbook(io.BytesIO(r.content))["data"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    # Ô gộp chỉ giữ giá trị ở dòng đầu; các dòng sau của nhóm để trống.
    top = next(i for i, row in enumerate(rows, start=2) if row[1] == "0314892345")
    assert rows[top - 2][0] == "Mã số thuế"
    assert rows[top - 1][:2] == [None, None]
    assert {rows[top - 2][3], rows[top - 1][3]} == {a_id, b_id}
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {f"A{top}:A{top + 1}", f"B{top}:B{top + 1}"} <= merged


# ── evidence_checks append-only ────────────────────────────────────────────────────────────
@pytest.mark.parametrize(
    "sql",
    ["UPDATE evidence_checks SET note = 'x'", "DELETE FROM evidence_checks"],
)
async def test_evidence_checks_are_append_only(
    api_client: AsyncClient, db_session: AsyncSession, company_id: uuid.UUID, sql: str
) -> None:
    await login_admin(api_client, db_session)
    await api_client.post(
        checks_url(company_id), json={"check_type": "email_domain", "result": "match"}
    )
    assert (await db_session.scalars(select(EvidenceCheck))).first() is not None
    with pytest.raises(DBAPIError, match="append-only"):
        await db_session.execute(text(sql))
