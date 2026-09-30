"""Xuất / nhập Excel loại bằng chứng + luật bằng chứng, và sửa luật (C6). Dữ liệu SYNTHETIC."""

import io
import uuid
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.audit import AuditLog
from app.core.spreadsheet import XLSX
from app.modules.auth.service import create_admin
from app.modules.companies.tests.helpers import PASSWORD, login_as
from app.modules.verification.models import EvidenceType, RequiredEvidenceRule
from app.modules.verification.tests.helpers import add_rule, add_type

pytestmark = pytest.mark.usefixtures("hs_seeded")

TYPES = "/api/admin/evidence-types"
RULES = "/api/admin/evidence-rules"
NIL = "00000000-0000-0000-0000-000000000000"
T_HEADER = ["code", "name_vi", "name_en", "group", "validity_months", "is_active"]
R_HEADER = ["category", "evidence_type_code", "is_required", "note"]


@pytest.fixture
async def admin(api_client: AsyncClient, db_session: AsyncSession) -> AsyncClient:
    await create_admin(db_session, "admin@evfta.eu", PASSWORD)
    r = await api_client.post(
        "/api/auth/login", json={"email": "admin@evfta.eu", "password": PASSWORD}
    )
    assert r.status_code == 200, r.text
    return api_client


def xlsx(header: list[str], *rows: list[Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "data"
    ws.append(header)
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def upload(
    client: AsyncClient, path: str, data: bytes, dry_run: bool = False
) -> dict[str, Any]:
    r = await client.post(
        f"{path}/import", params={"dry_run": dry_run}, files={"file": ("f.xlsx", data, XLSX)}
    )
    assert r.status_code == 200, r.text
    body: dict[str, Any] = r.json()
    return body


async def actions(session: AsyncSession) -> list[str]:
    return list(await session.scalars(select(AuditLog.action_type).order_by(AuditLog.created_at)))


ROUTES = [
    ("GET", f"{TYPES}/template.xlsx"),
    ("GET", f"{TYPES}/export.xlsx"),
    ("POST", f"{TYPES}/import"),
    ("GET", f"{RULES}/template.xlsx"),
    ("GET", f"{RULES}/export.xlsx"),
    ("POST", f"{RULES}/import"),
    ("PATCH", f"{RULES}/{NIL}"),
]


@pytest.mark.parametrize(("method", "path"), ROUTES)
async def test_routes_401_without_session(api_client: AsyncClient, method: str, path: str) -> None:
    assert (await api_client.request(method, path)).status_code == 401


@pytest.mark.parametrize("role", ["exporter", "buyer"])
@pytest.mark.parametrize(("method", "path"), ROUTES)
async def test_routes_403_for_other_roles(
    api_client: AsyncClient, role: str, method: str, path: str
) -> None:
    await login_as(api_client, role, f"{role}@x.vn")
    assert (await api_client.request(method, path)).status_code == 403


@pytest.mark.parametrize("path", [TYPES, RULES])
async def test_template_has_data_and_guide_sheets(admin: AsyncClient, path: str) -> None:
    r = await admin.get(f"{path}/template.xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["data", "huong_dan"]
    assert "review_status" not in [c.value for c in wb["data"][1]]
    assert wb["data"].max_row == 1


async def test_type_import_creates_unreviewed_and_is_idempotent(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    file = xlsx(T_HEADER, ["iso_9001", "ISO 9001", "ISO 9001", "quality", 36, "true"])
    body = await upload(admin, TYPES, file)
    assert (body["created"], body["applied"]) == (1, True)
    row = await db_session.get(EvidenceType, "iso_9001")
    assert row is not None and (row.reviewed_by, row.validity_months) == (None, 36)
    assert "evidence_type.create" in await actions(db_session)
    again = await upload(admin, TYPES, file)
    assert (again["created"], again["updated"], again["unchanged"]) == (0, 0, 1)


async def test_type_import_change_unreviews_reviewed_row(
    admin: AsyncClient, db_session: AsyncSession, reviewer_id: uuid.UUID
) -> None:
    await add_type(db_session, reviewer_id, "iso_9001")
    file = xlsx(T_HEADER, ["iso_9001", "Tên mới", "New name", "quality", None, "true"])
    body = await upload(admin, TYPES, file)
    assert (body["created"], body["updated"]) == (0, 1)
    row = await db_session.get(EvidenceType, "iso_9001")
    await db_session.refresh(row)
    assert row is not None and (row.reviewed_by, row.name_vi) == (None, "Tên mới")
    assert "evidence_type.update" in await actions(db_session)


async def test_type_export_roundtrip_is_unchanged(
    admin: AsyncClient, db_session: AsyncSession, reviewer_id: uuid.UUID
) -> None:
    await add_type(db_session, reviewer_id, "iso_9001", validity=12)
    exported = (await admin.get(f"{TYPES}/export.xlsx")).content
    body = await upload(admin, TYPES, exported)
    assert (body["created"], body["updated"], body["unchanged"]) == (0, 0, 1)
    row = await db_session.get(EvidenceType, "iso_9001")
    assert row is not None and row.reviewed_by is not None


async def test_type_import_bad_row_writes_nothing(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    file = xlsx(
        T_HEADER,
        ["iso_9001", "ISO 9001", "ISO 9001", "quality", 36, "true"],
        ["Bad Code", "x", "x", "quality", None, "true"],
        ["iso_14001", "x", "x", "quality", 0, "true"],
    )
    body = await upload(admin, TYPES, file)
    assert body["applied"] is False and [e["row"] for e in body["errors"]] == [3, 4]
    assert await db_session.get(EvidenceType, "iso_9001") is None


async def test_rule_import_needs_existing_type_and_category(
    admin: AsyncClient, db_session: AsyncSession, reviewer_id: uuid.UUID
) -> None:
    await add_type(db_session, reviewer_id, "iso_9001")
    file = xlsx(
        R_HEADER,
        ["agriculture", "iso_9001", "true", "ghi chú"],
        ["agriculture", "no_such_type", "true", None],
        ["no_such_category", "iso_9001", "true", None],
    )
    body = await upload(admin, RULES, file)
    assert body["applied"] is False and [e["row"] for e in body["errors"]] == [3, 4]


async def test_rule_import_and_roundtrip(admin: AsyncClient, db_session: AsyncSession) -> None:
    await add_type(db_session, None, "iso_9001")
    file = xlsx(R_HEADER, ["agriculture", "iso_9001", "false", "chỉ nhắc"])
    body = await upload(admin, RULES, file)
    assert (body["created"], body["applied"]) == (1, True)
    rule = (await db_session.scalars(select(RequiredEvidenceRule))).one()
    assert (rule.reviewed_by, rule.is_required, rule.note) == (None, False, "chỉ nhắc")
    again = await upload(admin, RULES, (await admin.get(f"{RULES}/export.xlsx")).content)
    assert (again["created"], again["updated"], again["unchanged"]) == (0, 0, 1)


async def test_patch_rule_unreviews_and_audits(
    admin: AsyncClient, db_session: AsyncSession, reviewer_id: uuid.UUID
) -> None:
    await add_type(db_session, reviewer_id, "iso_9001")
    await add_rule(db_session, reviewer_id, "iso_9001")
    rule = (await db_session.scalars(select(RequiredEvidenceRule))).one()
    r = await admin.patch(f"{RULES}/{rule.id}", json={"is_required": False, "note": "nhắc"})
    assert r.status_code == 200, r.text
    assert (r.json()["is_required"], r.json()["reviewed_by"]) == (False, None)
    assert "evidence_rule.update" in await actions(db_session)


async def test_patch_rule_errors(
    admin: AsyncClient, db_session: AsyncSession, reviewer_id: uuid.UUID
) -> None:
    await add_type(db_session, reviewer_id, "iso_9001")
    await add_type(db_session, reviewer_id, "iso_14001")
    await add_rule(db_session, reviewer_id, "iso_9001")
    await add_rule(db_session, reviewer_id, "iso_14001")
    rule = (
        await db_session.scalars(
            select(RequiredEvidenceRule).where(
                RequiredEvidenceRule.evidence_type_code == "iso_9001"
            )
        )
    ).one()
    assert (await admin.patch(f"{RULES}/{NIL}", json={"note": "x"})).status_code == 404
    bad = await admin.patch(f"{RULES}/{rule.id}", json={"evidence_type_code": "missing"})
    assert bad.status_code == 422
    null = await admin.patch(f"{RULES}/{rule.id}", json={"is_required": None})
    assert null.status_code == 422
    dup = await admin.patch(f"{RULES}/{rule.id}", json={"evidence_type_code": "iso_14001"})
    assert dup.status_code == 409
