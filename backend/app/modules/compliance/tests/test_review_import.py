"""SPEC_compliance_data_20_codes §6.3, §9: nhập kết quả duyệt của luật sư thương mại (import-review).

File Excel mẫu được dựng trong test. Người duyệt phải là admin có năng lực is_legal_reviewer."""

import datetime as dt
import io
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.audit import AuditLog
from app.core.errors import AppError
from app.modules.auth import service as auth
from app.modules.companies.tests.helpers import PASSWORD
from app.modules.compliance.models import (
    ComplianceEvidenceRequirement,
    ComplianceEvidenceType,
    ComplianceReviewIssue,
    EvidenceTypeMapping,
    HsCodeCompliance,
    ProductSpecificRule,
)
from app.modules.compliance.review_import import (
    MAP_SHEET,
    REQ_SHEET,
    TYPE_SHEET,
    ReviewImportError,
    import_review,
)
from app.modules.compliance.seed import load_seed
from app.modules.compliance.tests.test_seed_loader import SEED_DIR, VERSION

SHRIMP = "03061792"
HEADER = ["Mã CN", "Mã bằng chứng", "Điều kiện áp dụng", "Kết luận", "Người duyệt", "Ghi chú"]
EMAIL = "luat-tm@evfta.eu"


def workbook(
    rows: list[list[Any]],
    types: list[list[Any]] | None = None,
    mappings: list[list[Any]] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = REQ_SHEET
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    if types is not None:
        sheet = wb.create_sheet(TYPE_SHEET)
        sheet.append(["Mã", "Kết luận", "Ghi chú"])
        for r in types:
            sheet.append(r)
    if mappings is not None:
        sheet = wb.create_sheet(MAP_SHEET)
        sheet.append(["Mã compliance", "Kết luận", "Ghi chú"])
        for r in mappings:
            sheet.append(r)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def row(code: str, cond: str, conclusion: str | None, note: str | None = None) -> list[Any]:
    return [SHRIMP, code, cond, conclusion, None, note]


@pytest.fixture
async def seeded(db_session: AsyncSession) -> AsyncSession:
    await load_seed(db_session, SEED_DIR, VERSION)
    return db_session


@pytest.fixture
async def reviewer(seeded: AsyncSession) -> str:
    await auth.create_admin(seeded, EMAIL, PASSWORD)
    await auth.set_legal_reviewer(seeded, EMAIL, True)
    return EMAIL


async def _req(session: AsyncSession, code: str, cond: str) -> ComplianceEvidenceRequirement:
    found = await session.scalar(
        select(ComplianceEvidenceRequirement).where(
            ComplianceEvidenceRequirement.hs_code == SHRIMP,
            ComplianceEvidenceRequirement.evidence_type == code,
            ComplianceEvidenceRequirement.condition == cond,
        )
    )
    assert found is not None
    await session.refresh(found)
    return found


# --- quyền người duyệt -----------------------------------------------------------------------------


async def test_reviewer_must_be_a_legal_reviewer_admin(seeded: AsyncSession) -> None:
    data = workbook([row("BOM_ORIGIN", "ALWAYS", "DONG_Y")])
    with pytest.raises(AppError) as unknown:
        await import_review(seeded, data, "nobody@evfta.eu")
    assert unknown.value.status_code == 404
    await auth.create_admin(seeded, "plain-admin@evfta.eu", PASSWORD)
    with pytest.raises(AppError) as plain:
        await import_review(seeded, data, "plain-admin@evfta.eu")
    assert plain.value.status_code == 403
    assert (await _req(seeded, "BOM_ORIGIN", "ALWAYS")).reviewed_by is None


async def test_legal_reviewer_flag_is_admin_only_and_audited(seeded: AsyncSession) -> None:
    admin_id = await auth.create_admin(seeded, EMAIL, PASSWORD)
    assert await auth.is_legal_reviewer(seeded, admin_id) is False
    await auth.set_legal_reviewer(seeded, EMAIL, True)
    assert await auth.is_legal_reviewer(seeded, admin_id) is True
    audit = await seeded.scalar(
        select(AuditLog).where(AuditLog.action_type == "user.set_legal_reviewer")
    )
    assert audit is not None and audit.after_state == {"is_legal_reviewer": True}
    with pytest.raises(AppError):  # tài khoản không phải admin
        await auth.set_legal_reviewer(seeded, "exporter@x.vn", True)


# --- DONG_Y / BO / SUA ----------------------------------------------------------------------------


async def test_dong_y_reviews_rows_and_writes_audit(seeded: AsyncSession, reviewer: str) -> None:
    report = await import_review(
        seeded, workbook([row("BOM_ORIGIN", "ALWAYS", "DONG_Y")]), reviewer
    )
    assert report.approved == 1
    req = await _req(seeded, "BOM_ORIGIN", "ALWAYS")
    assert req.reviewed_by is not None and req.reviewed_at is not None
    audit = await seeded.scalar(
        select(AuditLog).where(
            AuditLog.entity_id == str(req.id),
            AuditLog.action_type == "compliance_evidence_requirement.review",
        )
    )
    assert audit is not None and audit.actor_id == req.reviewed_by
    assert audit.before_state is not None and audit.before_state["reviewed_by"] is None
    assert audit.after_state is not None and audit.after_state["reviewed_by"] == str(
        req.reviewed_by
    )


async def test_rerun_does_not_review_twice_or_double_audit(
    seeded: AsyncSession, reviewer: str
) -> None:
    data = workbook([row("BOM_ORIGIN", "ALWAYS", "DONG_Y")])
    await import_review(seeded, data, reviewer)
    again = await import_review(seeded, data, reviewer)
    assert again.approved == 0 and again.already_reviewed == 1
    audits = list(
        await seeded.scalars(
            select(AuditLog).where(AuditLog.action_type == "compliance_evidence_requirement.review")
        )
    )
    assert len(audits) == 1


async def test_bo_ends_validity_today(seeded: AsyncSession, reviewer: str) -> None:
    report = await import_review(seeded, workbook([row("BOM_ORIGIN", "ALWAYS", "BO")]), reviewer)
    assert report.ended == 1
    req = await _req(seeded, "BOM_ORIGIN", "ALWAYS")
    assert req.valid_until == dt.datetime.now(dt.UTC).date()
    assert req.reviewed_by is None  # BO không đồng nghĩa duyệt
    again = await import_review(seeded, workbook([row("BOM_ORIGIN", "ALWAYS", "BO")]), reviewer)
    assert again.ended == 0  # dòng đã hết hiệu lực không còn khớp


async def test_sua_goes_to_admin_queue_and_is_not_reviewed(
    seeded: AsyncSession, reviewer: str
) -> None:
    data = workbook([row("EUR1", "CONSIGNMENT_GT_6000", "SUA", "Điều 15.2(a): cần dẫn chiếu thêm")])
    report = await import_review(seeded, data, reviewer)
    assert report.issues_created == 1 and report.approved == 0
    req = await _req(seeded, "EUR1", "CONSIGNMENT_GT_6000")
    assert req.reviewed_by is None
    (issue,) = (await seeded.scalars(select(ComplianceReviewIssue))).all()
    assert issue.entity_id == str(req.id) and "dẫn chiếu" in issue.note
    again = await import_review(seeded, data, reviewer)
    assert again.issues_created == 0 and again.issues_existing == 1


async def test_blank_conclusion_is_skipped(seeded: AsyncSession, reviewer: str) -> None:
    report = await import_review(seeded, workbook([row("BOM_ORIGIN", "ALWAYS", None)]), reviewer)
    assert report.approved == report.ended == report.issues_created == 0
    assert (await _req(seeded, "BOM_ORIGIN", "ALWAYS")).reviewed_by is None


async def test_invalid_file_writes_nothing_and_reports_all_errors(
    seeded: AsyncSession, reviewer: str
) -> None:
    data = workbook(
        [
            row("BOM_ORIGIN", "ALWAYS", "DONG_Y"),  # hợp lệ nhưng không được ghi
            row("PURCHASE_DOCS", "ALWAYS", "TUY_Y"),
            row("EUR1", "CONSIGNMENT_GT_6000", "SUA"),  # thiếu ghi chú
            ["99999999", "EUR1", "ALWAYS", "DONG_Y", None, None],  # không có dòng
        ]
    )
    with pytest.raises(ReviewImportError) as exc:
        await import_review(seeded, data, reviewer)
    assert len(exc.value.errors) == 3
    assert (await _req(seeded, "BOM_ORIGIN", "ALWAYS")).reviewed_by is None
    assert (await seeded.scalars(select(ComplianceReviewIssue))).all() == []


async def test_not_an_xlsx_is_rejected(seeded: AsyncSession, reviewer: str) -> None:
    with pytest.raises(AppError):
        await import_review(seeded, b"not an excel file", reviewer)


async def test_the_shipped_review_workbook_is_inert_until_the_lawyer_fills_it(
    seeded: AsyncSession, reviewer: str
) -> None:
    data = (SEED_DIR / "evidence_matrix_review.xlsx").read_bytes()
    report = await import_review(seeded, data, reviewer)
    assert report.approved == report.ended == report.issues_created == 0


# --- sheet tuỳ chọn: loại bằng chứng và ánh xạ ------------------------------------------------------


async def test_optional_sheets_review_types_and_mappings(
    seeded: AsyncSession, reviewer: str
) -> None:
    data = workbook(
        [],
        types=[["EUR1", "DONG_Y", None], ["FARM_RECORD", "BO", None]],
        mappings=[["EUR1_ISSUED_12M", "DONG_Y", None]],
    )
    report = await import_review(seeded, data, reviewer)
    assert (report.approved, report.ended) == (2, 1)
    assert (await seeded.get(ComplianceEvidenceType, "EUR1")).reviewed_by is not None  # type: ignore[union-attr]
    mapping = await seeded.get(EvidenceTypeMapping, "EUR1_ISSUED_12M")
    assert mapping is not None and mapping.reviewed_by is not None
    bad = workbook([], mappings=[["EUR1_ISSUED_12M", "BO", None]])
    with pytest.raises(ReviewImportError):
        await import_review(seeded, bad, reviewer)


# --- §9: sau khi nhập file duyệt, mã đủ thành phần chuyển REVIEWED -------------------------------------


async def test_fully_reviewed_code_loses_the_disclaimer(
    api_client: AsyncClient, seeded: AsyncSession, reviewer: str
) -> None:
    body = {
        "hs_code": SHRIMP,
        "consignment_value_eur": "5000",
        "transit_third_country": False,
        "sourcing": "FARMED_IN_VN",
        "raw_material_source": "AQUACULTURE",
    }
    before = (await api_client.post("/api/public/origin", json=body)).json()
    assert before["review_state"] == "UNREVIEWED" and before["disclaimer"] is not None

    reqs = list(
        await seeded.scalars(
            select(ComplianceEvidenceRequirement).where(
                ComplianceEvidenceRequirement.hs_code == SHRIMP
            )
        )
    )
    rows = [row(r.evidence_type, r.condition.value, "DONG_Y") for r in reqs]
    await import_review(seeded, workbook(rows), reviewer)
    # Quy tắc xuất xứ và ánh xạ mã CN được duyệt qua các cổng duyệt khác (admin / đối chiếu CN 2026).
    admin = await auth.get_admin_by_email(seeded, reviewer)
    assert admin is not None
    rule = await seeded.scalar(
        select(ProductSpecificRule).where(ProductSpecificRule.hs_code == SHRIMP)
    )
    cn = await seeded.get(HsCodeCompliance, SHRIMP)
    assert rule is not None and cn is not None
    rule.reviewed_by, rule.reviewed_at = admin.id, dt.datetime.now(dt.UTC)
    cn.cn_mapping_verified = True
    await seeded.flush()

    after = (await api_client.post("/api/public/origin", json=body)).json()
    assert after["review_state"] == "REVIEWED" and after["disclaimer"] is None
    assert after["unreviewed_components"] == []
    assert after["required_evidence"]["review_state"] == "REVIEWED"


# --- hàng đợi admin (API) -----------------------------------------------------------------------------


async def test_admin_queue_api_lists_and_resolves(
    api_client: AsyncClient, seeded: AsyncSession, reviewer: str
) -> None:
    await import_review(
        seeded,
        workbook([row("EUR1", "CONSIGNMENT_GT_6000", "SUA", "Cần sửa căn cứ")]),
        reviewer,
    )
    url = "/api/admin/compliance-review-issues"
    assert (await api_client.get(url)).status_code == 401
    res = await api_client.post("/api/auth/login", json={"email": EMAIL, "password": PASSWORD})
    assert res.status_code == 200, res.text
    issues = (await api_client.get(url)).json()
    assert len(issues) == 1 and issues[0]["note"] == "Cần sửa căn cứ"
    done = await api_client.post(f"{url}/{issues[0]['id']}/resolve")
    assert done.status_code == 200 and done.json()["resolved_at"] is not None
    assert (await api_client.get(url)).json() == []
    assert len((await api_client.get(url, params={"only_open": "false"})).json()) == 1
