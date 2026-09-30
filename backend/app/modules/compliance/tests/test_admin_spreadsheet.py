"""Xuất / nhập Excel dòng thuế và PSR (admin). Dữ liệu là SYNTHETIC."""

import datetime as dt
import io
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.audit import AuditLog
from app.core.spreadsheet import XLSX
from app.modules.auth.service import create_admin
from app.modules.companies.tests.helpers import PASSWORD, login_as
from app.modules.compliance.models import ProductSpecificRule, TariffLine

pytestmark = pytest.mark.usefixtures("hs_seeded")

TARIFF = "/api/admin/tariff-lines"
RULES = "/api/admin/roo-rules"
FROM = (dt.datetime.now(dt.UTC).date() - dt.timedelta(days=30)).isoformat()
T_HEADER = ["hs_code", "destination", "duty_type", "mfn_rate", "evfta_rate_current", "valid_from"]
R_HEADER = ["hs_code", "rule_type", "threshold_pct", "requires_expert", "valid_from"]


def tariff_body(**over: Any) -> dict[str, Any]:
    body: dict[str, Any] = {
        "hs_code": "090121",
        "destination": "EU",
        "duty_type": "ad_valorem",
        "mfn_rate": "7.5",
        "valid_from": FROM,
    }
    body.update(over)
    return body


@pytest.fixture
async def admin(api_client: AsyncClient, db_session: AsyncSession) -> AsyncClient:
    await create_admin(db_session, "admin@evfta.eu", PASSWORD)
    r = await api_client.post(
        "/api/auth/login", json={"email": "admin@evfta.eu", "password": PASSWORD}
    )
    assert r.status_code == 200, r.text
    return api_client


def xlsx(header: list[str], *rows: list[Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "data"
    ws.append(header)
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def upload(
    client: AsyncClient, path: str, data: bytes, dry_run: bool = False
) -> dict[str, Any]:
    r = await client.post(
        f"{path}/import", params={"dry_run": dry_run}, files={"file": ("f.xlsx", data, XLSX)}
    )
    assert r.status_code == 200, r.text
    body: dict[str, Any] = r.json()
    return body


async def count(session: AsyncSession, model: Any) -> int:
    return int(await session.scalar(select(func.count()).select_from(model)) or 0)


async def actions(session: AsyncSession) -> list[str]:
    return list(await session.scalars(select(AuditLog.action_type).order_by(AuditLog.created_at)))


ROUTES = [
    ("GET", f"{TARIFF}/template.xlsx"),
    ("GET", f"{TARIFF}/export.xlsx"),
    ("POST", f"{TARIFF}/import"),
    ("GET", f"{RULES}/template.xlsx"),
    ("GET", f"{RULES}/export.xlsx"),
    ("POST", f"{RULES}/import"),
]


@pytest.mark.parametrize(("method", "path"), ROUTES)
async def test_spreadsheet_routes_401_without_session(
    api_client: AsyncClient, method: str, path: str
) -> None:
    assert (await api_client.request(method, path)).status_code == 401


@pytest.mark.parametrize("role", ["exporter", "buyer"])
@pytest.mark.parametrize(("method", "path"), ROUTES)
async def test_spreadsheet_routes_403_for_other_roles(
    api_client: AsyncClient, role: str, method: str, path: str
) -> None:
    await login_as(api_client, role, f"{role}@x.vn")
    assert (await api_client.request(method, path)).status_code == 403


@pytest.mark.parametrize(("path", "enum_column"), [(TARIFF, "duty_type"), (RULES, "rule_type")])
async def test_template_has_data_and_guide_sheets(
    admin: AsyncClient, path: str, enum_column: str
) -> None:
    r = await admin.get(f"{path}/template.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["data", "huong_dan"]
    header = [c.value for c in wb["data"][1]]
    assert header[0] == "hs_code"
    assert "review_status" not in header
    assert wb["data"].max_row == 1  # mẫu không chứa dữ liệu
    guide = {row[0].value: row for row in wb["huong_dan"].iter_rows(min_row=2)}
    assert guide[enum_column][3].value  # liệt kê giá trị hợp lệ của cột enum


async def test_import_creates_unreviewed_rows_with_audit(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    file = xlsx(T_HEADER, ["090121", "EU", "ad_valorem", 7.5, 0, FROM])
    body = await upload(admin, TARIFF, file)
    assert (body["created"], body["updated"], body["unchanged"]) == (1, 0, 0)
    assert body["applied"] is True and body["errors"] == []
    row = (await db_session.scalars(select(TariffLine))).one()
    assert (row.reviewed_by, str(row.mfn_rate)) == (None, "7.5000")
    assert "tariff_line.create" in await actions(db_session)


async def test_dry_run_writes_nothing(admin: AsyncClient, db_session: AsyncSession) -> None:
    file = xlsx(T_HEADER, ["090121", "EU", "ad_valorem", 7.5, 0, FROM])
    body = await upload(admin, TARIFF, file, dry_run=True)
    assert (body["created"], body["applied"], body["dry_run"]) == (1, False, True)
    assert await count(db_session, TariffLine) == 0
    assert "tariff_line.create" not in await actions(db_session)


async def test_reviewed_columns_in_file_cannot_approve(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    header = [*T_HEADER, "reviewed_by", "reviewed_at", "review_status"]
    file = xlsx(header, ["090121", "EU", "ad_valorem", 7.5, 0, FROM, "x", "2026-01-01", "reviewed"])
    await upload(admin, TARIFF, file)
    assert (await db_session.scalars(select(TariffLine))).one().reviewed_by is None


async def test_export_then_import_is_all_unchanged_and_keeps_review(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    line = (await admin.post(TARIFF, json=tariff_body(evfta_rate_current="0"))).json()
    assert (await admin.post(f"{TARIFF}/{line['id']}/review")).status_code == 200
    exported = (await admin.get(f"{TARIFF}/export.xlsx")).content
    result = await upload(admin, TARIFF, exported)
    assert (result["created"], result["updated"], result["unchanged"]) == (0, 0, 1)
    assert (await db_session.scalars(select(TariffLine))).one().reviewed_by is not None


async def test_import_changed_reviewed_row_updates_and_unreviews(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    line = (await admin.post(TARIFF, json=tariff_body())).json()
    await admin.post(f"{TARIFF}/{line['id']}/review")
    file = xlsx(T_HEADER, ["090121", "EU", "ad_valorem", 9, 0, FROM])
    result = await upload(admin, TARIFF, file)
    assert (result["created"], result["updated"], result["unchanged"]) == (0, 1, 0)
    row = (await db_session.scalars(select(TariffLine))).one()
    assert (row.reviewed_by, str(row.mfn_rate)) == (None, "9.0000")
    assert "tariff_line.update" in await actions(db_session)


async def test_one_bad_row_rejects_whole_file_with_row_numbers(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    file = xlsx(
        T_HEADER,
        ["090121", "EU", "ad_valorem", 7.5, 0, FROM],
        ["090121", "VNM", "ad_valorem", 7.5, 0, "2026-02-01"],
        ["999999", "EU", "ad_valorem", 7.5, 0, "2026-03-01"],
    )
    body = await upload(admin, TARIFF, file)
    assert body["applied"] is False
    assert [e["row"] for e in body["errors"]] == [3, 4]
    assert await count(db_session, TariffLine) == 0


async def test_duplicate_key_inside_file_is_an_error(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    row = ["090121", "EU", "ad_valorem", 7.5, 0, FROM]
    body = await upload(admin, TARIFF, xlsx(T_HEADER, row, row))
    assert body["applied"] is False and body["errors"][0]["row"] == 3
    assert await count(db_session, TariffLine) == 0


async def test_excess_decimals_are_rejected_not_rounded(
    admin: AsyncClient, db_session: AsyncSession
) -> None:
    file = xlsx(T_HEADER, ["090121", "EU", "ad_valorem", 12.34567, 0, FROM])
    body = await upload(admin, TARIFF, file)
    assert body["applied"] is False  # 5 chữ số thập phân: báo lỗi, không làm tròn ngầm
    assert await count(db_session, TariffLine) == 0


async def test_invalid_files_are_rejected(admin: AsyncClient) -> None:
    for content in (b"not a zip", xlsx(["hs_code"], ["090121"])):
        r = await admin.post(f"{TARIFF}/import", files={"file": ("f.xlsx", content, XLSX)})
        assert r.status_code == 422, r.text
    big = await admin.post(
        f"{TARIFF}/import", files={"file": ("f.xlsx", b"0" * (2 * 1024 * 1024 + 1), XLSX)}
    )
    assert big.status_code == 413


async def test_export_neutralises_formula_injection(admin: AsyncClient) -> None:
    formula = '=HYPERLINK("http://x")'
    body = tariff_body(duty_type="specific", mfn_rate=None, mfn_specific=formula)
    assert (await admin.post(TARIFF, json=body)).status_code == 201
    wb = load_workbook(io.BytesIO((await admin.get(f"{TARIFF}/export.xlsx")).content))
    cell = next(c for c in wb["data"][2] if str(c.value).startswith("=HYPERLINK"))
    assert cell.data_type == "s"


async def test_psr_import_export_roundtrip(admin: AsyncClient, db_session: AsyncSession) -> None:
    file = xlsx(
        R_HEADER,
        ["090121", "MaxNOM", 70, "false", FROM],
        ["030617", "WO", None, "true", FROM],
    )
    body = await upload(admin, RULES, file)
    assert (body["created"], body["applied"]) == (2, True)
    rows = list(await db_session.scalars(select(ProductSpecificRule)))
    assert len(rows) == 2 and all(r.reviewed_by is None for r in rows)
    again = await upload(admin, RULES, (await admin.get(f"{RULES}/export.xlsx")).content)
    assert (again["created"], again["updated"], again["unchanged"]) == (0, 0, 2)


async def test_psr_threshold_rules_are_enforced_on_import(admin: AsyncClient) -> None:
    file = xlsx(R_HEADER, ["090121", "MaxNOM", None, "false", FROM])
    body = await upload(admin, RULES, file)
    assert body["applied"] is False and body["errors"][0]["row"] == 2
