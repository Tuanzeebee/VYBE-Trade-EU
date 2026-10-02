"""Nhập kết quả duyệt của luật sư thương mại từ file Excel (SPEC_compliance_data_20_codes §6.3).

    uv run python -m app.modules.compliance.seed import-review evidence_matrix_review.xlsx \\
        --reviewer-email <email>

- Người duyệt phải là admin có năng lực `is_legal_reviewer` (scripts.grant_legal_reviewer).
- Sheet `Ma tran bang chung`: cột Mã CN, Mã bằng chứng, Điều kiện áp dụng, Kết luận, Ghi chú.
  `DONG_Y` → đặt reviewed_by/reviewed_at; `SUA` → KHÔNG duyệt, ghi ghi chú vào hàng đợi admin
  (compliance_review_issues); `BO` → valid_until = hôm nay. Kết luận trống = chưa xem, bỏ qua.
- Sheet tuỳ chọn `Loai bang chung` (cột Mã, Kết luận, Ghi chú) và `Anh xa bang chung`
  (cột Mã compliance, Kết luận, Ghi chú; không có BO) dùng cùng quy ước.
- Kiểm TOÀN BỘ file trước, có lỗi thì báo hết và không ghi gì. Chạy lại không ghi đôi; dòng đã duyệt
  không bị duyệt lại. Mỗi dòng đổi trạng thái ghi audit_logs before/after.
- Cột Người duyệt / Ngày duyệt trong file bị BỎ QUA: người duyệt là tài khoản truyền vào lệnh.
"""

import datetime as dt
import io
import uuid
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.audit import record
from app.core.errors import AppError
from app.modules.auth import service as auth
from app.modules.auth.schemas import CurrentUser
from app.modules.compliance.models import (
    ComplianceEvidenceRequirement,
    ComplianceEvidenceType,
    ComplianceReviewIssue,
    EvidenceCondition,
    EvidenceTypeMapping,
)

REQ_SHEET = "Ma tran bang chung"
TYPE_SHEET = "Loai bang chung"
MAP_SHEET = "Anh xa bang chung"
CONCLUSIONS = ("DONG_Y", "SUA", "BO")

REQ_ENTITY = "compliance_evidence_requirement"
TYPE_ENTITY = "compliance_evidence_type"
MAP_ENTITY = "evidence_type_mapping"


class ReviewImportError(Exception):
    def __init__(self, errors: list[str]) -> None:
        super().__init__("; ".join(errors))
        self.errors = errors


@dataclass
class ReviewReport:
    approved: int = 0
    already_reviewed: int = 0
    ended: int = 0
    already_ended: int = 0
    issues_created: int = 0
    issues_existing: int = 0
    skipped_blank: int = 0
    by_sheet: dict[str, int] = field(default_factory=dict)


@dataclass
class _Item:
    sheet: str
    excel_row: int
    entity: str
    row: Any
    conclusion: str
    note: str | None
    label: str


def _text(value: Any) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def _sheet_rows(wb: Any, name: str) -> list[tuple[int, dict[str, str | None]]]:
    iterator = wb[name].iter_rows(values_only=True)
    header = next(iterator, None)
    if header is None:
        return []
    names = [_text(h) or "" for h in header]
    rows: list[tuple[int, dict[str, str | None]]] = []
    for number, values in enumerate(iterator, start=2):
        row = {names[i]: _text(v) for i, v in enumerate(values) if i < len(names) and names[i]}
        if any(row.values()):
            rows.append((number, row))
    return rows


def _conclusion(row: dict[str, str | None], where: str, errors: list[str]) -> str | None:
    raw = row.get("Kết luận")
    if raw is None:
        return None
    value = raw.upper()
    if value not in CONCLUSIONS:
        errors.append(f"{where}: kết luận '{raw}' không hợp lệ (DONG_Y / SUA / BO)")
        return None
    return value


def _snapshot(row: Any) -> dict[str, Any]:
    return {
        "reviewed_by": str(row.reviewed_by) if row.reviewed_by else None,
        "reviewed_at": row.reviewed_at.isoformat() if row.reviewed_at else None,
        "valid_until": row.valid_until.isoformat() if getattr(row, "valid_until", None) else None,
    }


def _open(workbook: bytes) -> Any:
    try:
        return load_workbook(io.BytesIO(workbook), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        raise AppError("invalid_file", "Not a valid .xlsx file", 422) from error


async def _requirement_rows(
    session: AsyncSession, cn: str, code: str, condition: EvidenceCondition
) -> list[ComplianceEvidenceRequirement]:
    return list(
        await session.scalars(
            select(ComplianceEvidenceRequirement).where(
                ComplianceEvidenceRequirement.hs_code == cn,
                ComplianceEvidenceRequirement.evidence_type == code,
                ComplianceEvidenceRequirement.condition == condition,
            )
        )
    )


async def _collect(session: AsyncSession, wb: Any, today: dt.date) -> list[_Item]:
    errors: list[str] = []
    items: list[_Item] = []
    sheets = set(wb.sheetnames)
    if REQ_SHEET not in sheets:
        raise ReviewImportError([f"Thiếu sheet '{REQ_SHEET}'"])

    for number, row in _sheet_rows(wb, REQ_SHEET):
        where = f"{REQ_SHEET} hàng {number}"
        conclusion = _conclusion(row, where, errors)
        if conclusion is None:
            continue
        cn, code, condition = (
            row.get("Mã CN"),
            row.get("Mã bằng chứng"),
            row.get("Điều kiện áp dụng"),
        )
        try:
            cond = EvidenceCondition(condition or "")
        except ValueError:
            errors.append(f"{where}: điều kiện '{condition}' không hợp lệ")
            continue
        if cn is None or code is None:
            errors.append(f"{where}: thiếu Mã CN hoặc Mã bằng chứng")
            continue
        found = await _requirement_rows(session, cn, code, cond)
        if not found:
            errors.append(f"{where}: không có dòng yêu cầu {cn} / {code} / {condition}")
            continue
        # DONG_Y/SUA chỉ áp cho dòng còn hiệu lực; BO khớp cả dòng đã kết thúc (chạy lại không lỗi).
        for req in found:
            ended = req.valid_until is not None and req.valid_until <= today
            if ended and conclusion != "BO":
                continue
            items.append(
                _Item(
                    REQ_SHEET,
                    number,
                    REQ_ENTITY,
                    req,
                    conclusion,
                    row.get("Ghi chú"),
                    f"{cn} · {code} · {condition}",
                )
            )

    if TYPE_SHEET in sheets:
        for number, row in _sheet_rows(wb, TYPE_SHEET):
            where = f"{TYPE_SHEET} hàng {number}"
            conclusion = _conclusion(row, where, errors)
            if conclusion is None:
                continue
            type_row = await session.get(ComplianceEvidenceType, row.get("Mã") or "")
            if type_row is None:
                errors.append(f"{where}: không có loại bằng chứng '{row.get('Mã')}'")
                continue
            items.append(
                _Item(
                    TYPE_SHEET,
                    number,
                    TYPE_ENTITY,
                    type_row,
                    conclusion,
                    row.get("Ghi chú"),
                    f"loại {type_row.code}",
                )
            )

    if MAP_SHEET in sheets:
        for number, row in _sheet_rows(wb, MAP_SHEET):
            where = f"{MAP_SHEET} hàng {number}"
            conclusion = _conclusion(row, where, errors)
            if conclusion is None:
                continue
            if conclusion == "BO":
                errors.append(f"{where}: ánh xạ không có kết luận BO (hãy dùng SUA)")
                continue
            mapping = await session.get(EvidenceTypeMapping, row.get("Mã compliance") or "")
            if mapping is None:
                errors.append(f"{where}: không có ánh xạ '{row.get('Mã compliance')}'")
                continue
            items.append(
                _Item(
                    MAP_SHEET,
                    number,
                    MAP_ENTITY,
                    mapping,
                    conclusion,
                    row.get("Ghi chú"),
                    f"ánh xạ {mapping.compliance_code}",
                )
            )

    for item in items:
        if item.conclusion == "SUA" and not item.note:
            errors.append(f"{item.sheet} hàng {item.excel_row}: kết luận SUA phải có Ghi chú")
    if errors:
        raise ReviewImportError(errors)
    return items


def _entity_id(item: _Item) -> str:
    row = item.row
    if isinstance(row, ComplianceEvidenceRequirement):
        return str(row.id)
    if isinstance(row, ComplianceEvidenceType):
        return row.code
    return str(row.compliance_code)


async def _legal_reviewer(session: AsyncSession, email: str) -> CurrentUser:
    reviewer = await auth.get_admin_by_email(session, email)
    if reviewer is None:
        raise AppError("reviewer_not_found", "No admin account with this email", 404)
    if not await auth.is_legal_reviewer(session, reviewer.id):
        raise AppError("reviewer_not_legal", "This account is not a legal reviewer", 403)
    return reviewer


async def import_review(
    session: AsyncSession, workbook: bytes, reviewer_email: str
) -> ReviewReport:
    """Nhập kết quả duyệt. Không commit — người gọi commit (một transaction) hoặc rollback."""
    reviewer = await _legal_reviewer(session, reviewer_email)
    now = dt.datetime.now(dt.UTC)
    wb = _open(workbook)
    items = await _collect(session, wb, now.date())
    report = ReviewReport()
    for item in items:
        report.by_sheet[item.sheet] = report.by_sheet.get(item.sheet, 0) + 1
        if item.conclusion == "DONG_Y":
            await _approve(session, item, reviewer.id, now, report)
        elif item.conclusion == "BO":
            await _end(session, item, reviewer.id, now, report)
        else:
            await _issue(session, item, reviewer.id, report)
    await session.flush()
    return report


async def _approve(
    session: AsyncSession,
    item: _Item,
    reviewer_id: uuid.UUID,
    now: dt.datetime,
    report: ReviewReport,
) -> None:
    row = item.row
    if row.reviewed_by is not None:
        report.already_reviewed += 1
        return
    before = _snapshot(row)
    row.reviewed_by, row.reviewed_at = reviewer_id, now
    report.approved += 1
    await record(
        session,
        actor_id=reviewer_id,
        action_type=f"{item.entity}.review",
        entity_type=item.entity,
        entity_id=_entity_id(item),
        before=before,
        after=_snapshot(row),
    )


async def _end(
    session: AsyncSession,
    item: _Item,
    reviewer_id: uuid.UUID,
    now: dt.datetime,
    report: ReviewReport,
) -> None:
    row = item.row
    today = now.date()
    if row.valid_until is not None and row.valid_until <= today:
        report.already_ended += 1
        return
    before = _snapshot(row)
    row.valid_until = max(today, row.valid_from + dt.timedelta(days=1))
    report.ended += 1
    await record(
        session,
        actor_id=reviewer_id,
        action_type=f"{item.entity}.end",
        entity_type=item.entity,
        entity_id=_entity_id(item),
        before=before,
        after=_snapshot(row),
    )


async def _issue(
    session: AsyncSession, item: _Item, reviewer_id: uuid.UUID, report: ReviewReport
) -> None:
    note = item.note or ""
    existing = await session.scalar(
        select(ComplianceReviewIssue.id).where(
            ComplianceReviewIssue.entity_type == item.entity,
            ComplianceReviewIssue.entity_id == _entity_id(item),
            ComplianceReviewIssue.note == note,
            ComplianceReviewIssue.resolved_at.is_(None),
        )
    )
    if existing is not None:
        report.issues_existing += 1
        return
    issue = ComplianceReviewIssue(
        entity_type=item.entity,
        entity_id=_entity_id(item),
        label=item.label[:255],
        note=note,
        created_by=reviewer_id,
    )
    session.add(issue)
    await session.flush()
    report.issues_created += 1
    await record(
        session,
        actor_id=reviewer_id,
        action_type="compliance_review_issue.create",
        entity_type="compliance_review_issue",
        entity_id=str(issue.id),
        before=None,
        after={"entity_type": item.entity, "entity_id": issue.entity_id, "note": note},
    )


# --- hàng đợi admin ---


async def list_issues(session: AsyncSession, only_open: bool = True) -> list[ComplianceReviewIssue]:
    query = select(ComplianceReviewIssue).order_by(ComplianceReviewIssue.created_at.desc())
    if only_open:
        query = query.where(ComplianceReviewIssue.resolved_at.is_(None))
    return list(await session.scalars(query))


async def resolve_issue(
    session: AsyncSession, actor: CurrentUser, issue_id: uuid.UUID
) -> ComplianceReviewIssue:
    issue = await session.get(ComplianceReviewIssue, issue_id)
    if issue is None:
        raise AppError("issue_not_found", "Review issue not found", 404)
    if issue.resolved_at is None:
        issue.resolved_by, issue.resolved_at = actor.id, dt.datetime.now(dt.UTC)
        await record(
            session,
            actor_id=actor.id,
            action_type="compliance_review_issue.resolve",
            entity_type="compliance_review_issue",
            entity_id=str(issue.id),
            before={"resolved": False},
            after={"resolved": True},
        )
        await session.commit()
    return issue
